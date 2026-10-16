#!/usr/bin/env python3
"""Build public/course-template.xlsx, the sample spreadsheet users can fill in and import."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "Courses"
ws.append(["Course", "Title", "Credits", "Grade", "Term", "Status"])
for row in [
    ["POLS 1101", "The American Political System", 3, "A-", "Fall 2025", "completed"],
    ["PHIL 1066", "Introduction to Logic", 3, "B+", "Fall 2025", "completed"],
    ["ENG 0802", "Analytical Reading and Writing", 4, "A", "Spring 2026", "completed"],
    ["POLS 1201", "Foreign Governments and Politics", 3, "", "Fall 2026", "in progress"],
    ["CSI 2111", "Argumentation and Advocacy", 3, "", "Spring 2027", "planned"],
]:
    ws.append(row)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="9D2235")
for col, w in zip("ABCDEF", (12, 38, 9, 8, 14, 14)):
    ws.column_dimensions[col].width = w
dv = DataValidation(type="list", formula1='"completed,in progress,planned"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("F2:F500")
ws.freeze_panes = "A2"

help_ws = wb.create_sheet("How to use")
for line in [
    ["Temple Major Tracker: course import template"],
    [""],
    ["Course: subject + number, e.g. POLS 1101 (required). You may instead use separate 'Subject' and 'Number' columns."],
    ["Credits: semester hours. If blank, the bulletin's credit value is used."],
    ["Grade: letter grade (A, A-, B+ ... F), P, CR, TR or W. Leave blank for courses not yet graded."],
    ["Status: completed, in progress, or planned. If blank, a graded course is treated as completed."],
    ["Delete the example rows, add your own, save as .xlsx or .csv, then drop the file on the My courses tab."],
]:
    help_ws.append(line)
help_ws.column_dimensions["A"].width = 110
help_ws["A1"].font = Font(bold=True, size=13)

out = Path(__file__).resolve().parent.parent / "public" / "course-template.xlsx"
wb.save(out)
print("wrote", out)
